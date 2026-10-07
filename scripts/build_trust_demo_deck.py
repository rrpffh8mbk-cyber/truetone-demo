"""Six-slide editable deck and matching browser-printable PDF source."""
from pathlib import Path
import json,html
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'docs/forensics';summary=json.loads((ROOT/'data/forensics/test-summary.json').read_text())
slides=[
 {'kicker':'试色有谱 · TrueTone','title':'先核验这条试色，\n再决定要不要信。','subtitle':'信任守护师场景原型 · 按团队提供的赛题方向调整','body':['同一支口红的网上展示可能因为光照、版本、原生唇色或局部编辑而不同。','我们的入口是核验信息：发现可解释的线索、对照来源，再决定哪些证据值得参考。','自拍颜色预览放在最后，作为辅助体验。'],'notes':'演示以取证开始。不声称判断作者真假，也不将颜色差异等同于篡改。赛题方向来自团队提供的说明，未独立验证官方规则。'},
 {'kicker':'核心机制','title':'一个 Evidence / Trust Agent，\n编排五个核验步骤。','body':['① 图像处理线索：找可比源图、校正整体曝光、定位局部残差。','② 图文核对：对照画面可见的颜色与浓淡；无法验证的自述保留待核验。','③ 同产品证据：同色号、同产品线与版本；保留具体负面体验。','④ 个人相关性：质量与已确认标签匹配；未知条件不补猜。','⑤ 下一步建议：索取原图、确认版本、暂缓作为颜色依据。'],'notes':'当前 Agent 是确定性编排器，不是四个独立 LLM 投票。源图定位独立于唇部颜色相似度与官方色卡差异。实验性预训练模型只给全图响应，不能提供 AI 区域。'},
 {'kicker':'来源隔离的受控实验','title':'先留出测试原帖，\n再冻结检测阈值。','body':['24 个原帖组，四个色号各 6 组；12 组调参、12 组保留测试。','每组包含 4 类正常处理对照与 5 类局部编辑；拼接供体不跨分组。','区域答案来自实际编辑 alpha；检测器不读取操作标签或区域答案。','6 张已知生成图来自一次 image_gen 生成，全部留作测试；仅一个生成家族。','对应基准图是核验输入。这不是缺少源图的盲测；对照来源也未经真实性认证。'],'notes':'108 个调参样本。保留测试 108 个源图派生样本加 6 张 AI 生成图。配置只由调参组选择；生成阈值由对照最高响应加 0.05 确定。生成图分辨率调整未改变阈值，冻结记录与清单修订可复核。'},
 {'kicker':'公开成功，也公开失败','title':'能定位部分局部编辑，\n目前还不能通用识别 AI 图。','body':['局部检测：40 / 60 新增编辑被检出；48 个正常处理对照未误报。','平均定位 IoU 0.399；漏检与无法匹配按 0 计入。','较轻局部改色：0 / 12。已知 AI 生成图触发模型：0 / 6。','因此生成模型保留实验入口，不能作为真实性通过门槛。'],'image':'localization-example.png','notes':'四列依次为比较源图、受控拼接、已知区域和实际检测。源图自身可能已处理。观察到零误报不等于真实误报率为零。6 张生成图相关，不能声称多生成器覆盖或模型泛化准确率。'},
 {'kicker':'正式演示顺序','title':'图像可疑区域 → 图文承诺 →\n产品版本 → 购买参考','body':['先打开篡改风险测试台，比较对照、改色与拼接；解释橙色区域和无法核验。','再检查种草案例：绝对承诺不能靠相近颜色获得“真实认证”。','用兰蔻 274 展示同号不同产品线的边界；版本不明确时不混算。','最后看条件匹配的 TOP3 图与具体体验；预览从真实证据中取色，辅助理解。'],'notes':'不把虚拟试妆放在开场。点击每个示例后实际运行检测，结果不读取答案。生成模型失败要主动说，不回避。网址：https://rrpffh8mbk-cyber.github.io/truetone-demo/forensics.html'},
 {'kicker':'品牌闭环与下一步','title':'把核验结果，\n变成可执行的素材改进。','body':['消费者：暂缓可疑颜色依据 → 索取源图 → 参考同产品、相近条件的真实体验。','品牌：回收需核验的素材 → 补拍同光照标准图 → 标清版本 → 调整详情页证据。','待验证指标：素材核验完成率、版本澄清率、证据点击与购买决策耗时。','下一步：多生成器、真实编辑过程、外部人工盲评与来源级独立测试。','边界：目前四个色号；未完成通用伪造、刷评或虚构文本识别，也未证明转化提升。'],'notes':'品牌收益是可测试的产品机制，而不是已证实的商业提升。保留颜色、证据和个人匹配的解释边界。让评委看到真实技术和未解决的问题。'}
]
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
BG='15110F';FG='F6EFE6';GOLD='D8A76D';MUTED='C5B5A5'
def box(slide,text,x,y,w,h,size=20,color=FG,bold=False):
 shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));frame=shape.text_frame;frame.word_wrap=True;frame.margin_left=frame.margin_right=0
 for i,line in enumerate(text.split('\n')):
  p=frame.paragraphs[0] if i==0 else frame.add_paragraph();p.text=line;p.font.name='Noto Sans CJK SC';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(10)
 return shape
for i,d in enumerate(slides):
 slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(BG)
 box(slide,d['kicker'],.65,.35,12,.5,16,GOLD);box(slide,d['title'],.65,1,12,1.5,32,FG,True)
 if d.get('subtitle'):box(slide,d['subtitle'],.65,2.55,12,.5,19,GOLD)
 if d.get('image'):
  for j,line in enumerate(d['body']):box(slide,line,.65,2.55+j*.44,12,.43,18,MUTED)
  slide.shapes.add_picture(str(OUT/d['image']),Inches(.65),Inches(4.5),width=Inches(12))
 else:
  top=3.3 if d.get('subtitle') else 2.9
  for j,line in enumerate(d['body']):box(slide,line,.65,top+j*.64,12,.62,20,MUTED)
 box(slide,'试色有谱 · 受控实验原型 · 2026-10-07',.65,7.08,11,.3,11,MUTED);box(slide,f'{i+1:02d} / 06',12,7.08,.7,.3,11,GOLD);slide.notes_slide.notes_text_frame.text=d['notes']
prs.save(OUT/'TrueTone_Trust_Guardian_Demo.pptx')
style='''@page{size:13.333in 7.5in;margin:0}*{box-sizing:border-box}body{margin:0;background:#15110f;color:#f6efe6;font-family:"Noto Sans CJK SC",sans-serif}.slide{width:13.333in;height:7.5in;padding:.35in .65in;position:relative;break-after:page;overflow:hidden}.kicker{color:#d8a76d;font-size:16pt;margin:0 0 .4in}h1{font-size:32pt;line-height:1.35;margin:0 0 .35in;white-space:pre-line}.subtitle{color:#d8a76d;font-size:19pt}.body{margin-top:.3in}.body p{color:#c5b5a5;font-size:20pt;line-height:1.6;margin:0 0 .19in}.image-slide .body{margin-top:0}.image-slide .body p{font-size:18pt;margin:0 0 .05in}.image-slide img{width:12in;height:2.45in;object-fit:contain;margin-top:.12in}.foot{position:absolute;bottom:.25in;left:.65in;color:#c5b5a5;font-size:11pt}.num{position:absolute;right:.65in;bottom:.25in;color:#d8a76d;font-size:11pt}'''
parts=['<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>试色有谱 · 信任守护师演示</title><style>'+style+'</style><body>']
for i,d in enumerate(slides):
 parts.append('<section class="slide '+('image-slide' if d.get('image') else '')+'"><p class="kicker">'+html.escape(d['kicker'])+'</p><h1>'+html.escape(d['title'])+'</h1>'+('<p class="subtitle">'+html.escape(d['subtitle'])+'</p>' if d.get('subtitle') else '')+'<div class="body">'+''.join('<p>'+html.escape(t)+'</p>' for t in d['body'])+'</div>'+('<img src="'+d['image']+'" alt="源图、受控拼接、编辑区域和检测区域">' if d.get('image') else '')+'<span class="foot">试色有谱 · 受控实验原型 · 2026-10-07</span><span class="num">'+str(i+1)+' / 06</span></section>')
parts.append('</body></html>');(OUT/'TrueTone_Trust_Guardian_Demo.html').write_text('\n'.join(parts))
(OUT/'presentation-notes.md').write_text('# 新版六页演示稿讲述要点\n\n'+ '\n\n'.join(f"## {i+1}. {d['title'].replace(chr(10),'')}\n\n{d['notes']}" for i,d in enumerate(slides))+'\n')
print('Created six-page PPTX, matching HTML and speaking notes')
